"""Excel（店・駅・移動手段の3シート）を読み、データ（dict）に変換する。API 側の変換処理。"""

from io import BytesIO

from openpyxl import load_workbook

# シート名 → {Excel の列名: データのキー}
SHEETS = {
    "店": {
        "id": "id", "名前": "name", "エリア": "area", "ジャンル": "genre", "緯度": "lat", "経度": "lon",
        "開店": "open", "閉店": "close", "滞在時間(分)": "stay", "価格(円)": "price", "満足度": "satisfaction",
        "満腹度": "fullness", "出典URL": "source", "調べた日": "checked", "メモ": "memo",
    },
    "駅": {"id": "id", "名前": "name", "緯度": "lat", "経度": "lon"},
    "移動手段": {
        "手段": "name", "経路": "route", "速度(m/分)": "speed", "迂回係数": "detour", "端の徒歩(分)": "end_walk",
        "待ち時間(分)": "wait", "料金方式": "fare_type", "料金の値": "fare_value", "疲労係数": "fatigue",
        "出典URL": "source", "メモ": "memo",
    },
}
KEYS = {"店": "shops", "駅": "stations", "移動手段": "modes"}
# 空欄でもよい列
OPTIONAL = {"エリア", "ジャンル", "出典URL", "調べた日", "メモ", "料金の値"}


class ExcelError(ValueError):
    pass


def _cell(v):
    if v is None:
        return ""
    if hasattr(v, "strftime"):  # 時刻や日付のセル
        return v.strftime("%H:%M") if hasattr(v, "hour") and not hasattr(v, "year") else v.strftime("%Y-%m-%d")
    return v


def read_dataset(content: bytes) -> dict:
    try:
        wb = load_workbook(BytesIO(content), read_only=True, data_only=True)
    except Exception as e:
        raise ExcelError(f"Excel ファイルとして読めません（{e}）") from e

    data = {}
    for sheet, columns in SHEETS.items():
        if sheet not in wb.sheetnames:
            raise ExcelError(f"シート「{sheet}」がありません")
        rows = list(wb[sheet].iter_rows(values_only=True))
        if not rows:
            raise ExcelError(f"シート「{sheet}」が空です")
        header = [str(h).strip() if h is not None else "" for h in rows[0]]
        missing = [c for c in columns if c not in header]
        if missing:
            raise ExcelError(f"シート「{sheet}」に列がありません: {', '.join(missing)}")
        items = []
        for n, row in enumerate(rows[1:], start=2):
            if all(v is None or str(v).strip() == "" for v in row):
                continue
            item = {}
            for col, key in columns.items():
                v = _cell(row[header.index(col)] if header.index(col) < len(row) else None)
                if v == "" and col not in OPTIONAL:
                    raise ExcelError(f"シート「{sheet}」の {n} 行目「{col}」が空欄です")
                item[key] = v
            items.append(item)
        data[KEYS[sheet]] = items
    return data
